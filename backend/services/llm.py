import asyncio
import json
import logging
import os
import random
import re
import time
from contextlib import asynccontextmanager

from openai import APIStatusError, AsyncOpenAI

from schemas.api import GenerateRequest, GenerateResponse, HistoryMessage, UsageStats
from services.context_budget import (
    clamp_text_to_token_budget,
    estimate_tokens as estimate_context_tokens,
    select_recent_messages_by_token_budget,
)

logger = logging.getLogger("ai-orchestrator")

# Only OpenRouter is used as the single provider
OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY")

clients: dict[str, AsyncOpenAI] = {}

def _get_env_float(name: str, default: float) -> float:
    raw = os.getenv(name)
    if raw is None or not str(raw).strip():
        return default
    try:
        return float(raw)
    except (TypeError, ValueError):
        return default


API_TIMEOUT     = _get_env_float("LLM_API_TIMEOUT", 90.0)
API_MAX_RETRIES = 0
LLM_MAX_CONCURRENT_REQUESTS = max(1, int(os.getenv("LLM_MAX_CONCURRENT_REQUESTS", "24")))
LLM_CONCURRENCY_ACQUIRE_TIMEOUT_SECONDS = _get_env_float(
    "LLM_CONCURRENCY_ACQUIRE_TIMEOUT_SECONDS",
    15.0,
)
AUTO_CONTINUE_MAX_PASSES = max(0, int(os.getenv("LLM_AUTO_CONTINUE_MAX_PASSES", "6")))
CHAT_HISTORY_TOKEN_BUDGET = max(0, int(os.getenv("LLM_CHAT_HISTORY_TOKEN_BUDGET", "8000")))
TUTOR_HISTORY_TOKEN_BUDGET = max(0, int(os.getenv("LLM_TUTOR_HISTORY_TOKEN_BUDGET", "2500")))
INLINE_ATTACHMENT_TOKEN_BUDGET = max(0, int(os.getenv("LLM_INLINE_ATTACHMENT_TOKEN_BUDGET", "12000")))
CONTEXT_SAFETY_MARGIN_TOKENS = max(256, int(os.getenv("LLM_CONTEXT_SAFETY_MARGIN_TOKENS", "1024")))
CONTINUE_PROMPT = (
    "Continue exactly from where you stopped. "
    "Do not repeat any previous text. "
    "Keep the same language, formatting, and structure. "
    "Output only the continuation."
)
AUTO_CONTINUE_FINISH_REASONS = {"length", "max_tokens"}
_llm_concurrency_semaphore = asyncio.Semaphore(LLM_MAX_CONCURRENT_REQUESTS)


if OPENROUTER_API_KEY:
    clients["openrouter"] = AsyncOpenAI(
        api_key=OPENROUTER_API_KEY,
        base_url="https://openrouter.ai/api/v1",
        timeout=API_TIMEOUT,
        max_retries=API_MAX_RETRIES,
        default_headers={
            "HTTP-Referer": os.getenv("APP_URL", "http://localhost:3000"),
            "X-Title": os.getenv("APP_NAME", "Nexa"),
        },
    )
    logger.info("OpenRouter client initialized")

# ── Pricing (USD per 1M tokens) ─────────────────────────────────────────────
# [input_per_1m, output_per_1m]
MODEL_PRICING: dict[str, tuple[float, float]] = {
    "anthropic/claude-sonnet-4.5":         (3.00,  15.00),
    "anthropic/claude-sonnet-4-5":         (3.00,  15.00),
    "anthropic/claude-haiku-4.5":          (1.00,   5.00),
    "anthropic/claude-4.5-haiku-20251001": (1.00,   5.00),
    "anthropic/claude-haiku-4-5-20251001": (1.00,   5.00),
    "openai/gpt-4o":                       (2.50,  10.00),
    "openai/gpt-4o-mini":                  (0.15,   0.60),
    "openai/o4-mini":                      (1.10,   4.40),
    "google/gemini-2.5-flash":             (0.30,   2.50),
    "google/gemini-2.5-flash-preview":     (0.30,   2.50),
    "google/gemini-2.5-flash-preview-09-2025": (0.30,   2.50),
    "google/gemini-2.0-flash-001":         (0.10,   0.40),
    "google/gemini-2.5-pro":               (1.25,  10.00),
    "google/gemini-2.5-pro-preview":       (1.25,  10.00),
    "google/gemini-2.5-pro-preview-06-05": (1.25,  10.00),
}


def estimate_cost_usd(api_model: str, prompt_tokens: int, completion_tokens: int) -> float:
    pricing = MODEL_PRICING.get(api_model)
    if not pricing:
        return 0.0
    input_rate, output_rate = pricing
    return round((prompt_tokens * input_rate + completion_tokens * output_rate) / 1_000_000, 6)


def _api_name_candidates(model_info: dict) -> list[str]:
    primary = str(model_info.get("api_name", "")).strip()
    aliases = model_info.get("api_aliases") or []
    candidates: list[str] = []
    for value in [primary, *aliases]:
        candidate = str(value).strip()
        if candidate and candidate not in candidates:
            candidates.append(candidate)
    return candidates


def _is_invalid_model_error(exc: Exception) -> bool:
    if isinstance(exc, APIStatusError) and getattr(exc, "status_code", None) == 400:
        message = str(exc).lower()
        return "not a valid model id" in message or "invalid model" in message
    message = str(exc).lower()
    return "not a valid model id" in message or "invalid model" in message


def _effective_max_tokens(api_model: str, requested: int) -> int:
    # OpenAI o-series models may spend completion budget on internal reasoning.
    # Too small a budget can produce an empty visible answer via OpenRouter.
    if api_model == "openai/o4-mini":
        return max(requested, 512)
    return requested


@asynccontextmanager
async def acquire_llm_capacity(model: str, provider: str):
    acquired = False
    started_at = time.monotonic()
    try:
        await asyncio.wait_for(
            _llm_concurrency_semaphore.acquire(),
            timeout=LLM_CONCURRENCY_ACQUIRE_TIMEOUT_SECONDS,
        )
        acquired = True
    except asyncio.TimeoutError as exc:
        logger.warning(
            "[llm] capacity_timeout",
            extra={
                "provider": provider,
                "model": model,
                "max_concurrent_requests": LLM_MAX_CONCURRENT_REQUESTS,
                "acquire_timeout_seconds": LLM_CONCURRENCY_ACQUIRE_TIMEOUT_SECONDS,
            },
        )
        raise RuntimeError("llm_capacity_timeout") from exc

    wait_ms = int((time.monotonic() - started_at) * 1000)
    if wait_ms >= 250:
        logger.info(
            "[llm] capacity_wait",
            extra={
                "provider": provider,
                "model": model,
                "wait_ms": wait_ms,
            },
        )

    try:
        yield
    finally:
        if acquired:
            _llm_concurrency_semaphore.release()


# ── Model registry (all via OpenRouter) ────────────────────────────────────
# Paid models
AVAILABLE_MODELS: dict[str, dict] = {
    # Claude (Anthropic)
    "claude-sonnet-4-5":    {"provider": "openrouter", "label": "Claude Sonnet 4.5",    "api_name": "anthropic/claude-sonnet-4.5",         "api_aliases": ["anthropic/claude-sonnet-4-5"], "vision": True,  "free": False, "context": 1000000},
    "claude-haiku-4-5":     {"provider": "openrouter", "label": "Claude Haiku 4.5",     "api_name": "anthropic/claude-haiku-4.5",          "api_aliases": ["anthropic/claude-4.5-haiku-20251001", "anthropic/claude-haiku-4-5-20251001"], "vision": True,  "free": False, "context": 200000},
    # GPT (OpenAI)
    "gpt-4o":               {"provider": "openrouter", "label": "GPT-4o",               "api_name": "openai/gpt-4o",                       "vision": True,  "free": False, "context": 128000},
    "gpt-4o-mini":          {"provider": "openrouter", "label": "GPT-4o Mini",           "api_name": "openai/gpt-4o-mini",                  "vision": True,  "free": False, "context": 128000},
    "o4-mini":              {"provider": "openrouter", "label": "o4 Mini",               "api_name": "openai/o4-mini",                      "vision": True,  "free": False, "context": 128000},
    # Gemini (Google)
    "gemini-2.5-flash":     {"provider": "openrouter", "label": "Gemini 2.5 Flash",     "api_name": "google/gemini-2.5-flash",             "api_aliases": ["google/gemini-2.5-flash-preview", "google/gemini-2.5-flash-preview-09-2025"], "vision": True,  "free": False, "context": 1048576},
    "gemini-2.0-flash":     {"provider": "openrouter", "label": "Gemini 2.0 Flash",     "api_name": "google/gemini-2.0-flash-001",         "vision": True,  "free": False, "context": 1000000},
    "gemini-2.5-pro":       {"provider": "openrouter", "label": "Gemini 2.5 Pro",       "api_name": "google/gemini-2.5-pro",               "api_aliases": ["google/gemini-2.5-pro-preview", "google/gemini-2.5-pro-preview-06-05"], "vision": True,  "free": False, "context": 1048576},
}

# Helpers
def estimate_tokens(text: str) -> int:
    return estimate_context_tokens(text)


def _resolve_history_budget(
    *,
    model_context: int | None,
    reserved_tokens: int,
    default_budget: int = CHAT_HISTORY_TOKEN_BUDGET,
) -> int:
    if model_context is None or model_context <= 0:
        return default_budget
    available = model_context - reserved_tokens - CONTEXT_SAFETY_MARGIN_TOKENS
    if available <= 0:
        return 0
    return max(0, min(default_budget, available))


def _content_token_estimate(content: str | list | dict) -> int:
    if isinstance(content, str):
        return estimate_tokens(content)
    if isinstance(content, list):
        parts: list[str] = []
        for part in content:
            if isinstance(part, dict):
                if part.get("type") == "text":
                    parts.append(str(part.get("text", "")))
                elif part.get("type") == "image_url":
                    parts.append("[image attachment]")
            else:
                parts.append(str(part))
        return estimate_tokens("\n".join(parts))
    return estimate_tokens(str(content))


def get_client_for_model(model_id: str) -> tuple[AsyncOpenAI | None, dict | None]:
    model_info = AVAILABLE_MODELS.get(model_id)
    if not model_info:
        return None, None
    provider = model_info["provider"]
    client   = clients.get(provider)
    return client, model_info



_TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv",
    ".json", ".jsonl", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".env",
    ".xml", ".html", ".htm", ".svg",
    ".py", ".js", ".ts", ".tsx", ".jsx", ".mjs", ".cjs",
    ".java", ".kt", ".go", ".rs", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".rb", ".php", ".swift", ".scala", ".sh", ".bash", ".zsh",
    ".sql", ".graphql", ".proto", ".css", ".scss", ".less",
    ".log", ".diff", ".patch", ".tex", ".bib",
}


def _extract_inline_text(raw: bytes, mime: str, filename: str) -> str:
    """Extract readable text from raw bytes for injection into LLM context."""
    import base64 as _b64
    from pathlib import Path

    ext = Path(filename).suffix.lower()

    # PDF — use pymupdf
    if mime == "application/pdf" or ext == ".pdf":
        try:
            import fitz  # type: ignore[import]
            doc = fitz.open(stream=raw, filetype="pdf")
            text = "\n".join(page.get_text() for page in doc)
            doc.close()
            return text
        except ImportError:
            return "(PDF text extraction unavailable — install pymupdf)"
        except Exception as exc:
            return f"(PDF extraction failed: {exc})"

    # Word .docx
    if ext == ".docx" or "wordprocessingml" in mime:
        try:
            import docx, io  # type: ignore[import]
            doc = docx.Document(io.BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            return "(DOCX extraction unavailable — install python-docx)"
        except Exception as exc:
            return f"(DOCX extraction failed: {exc})"

    # Excel .xlsx
    if ext in (".xlsx", ".xls") or "spreadsheetml" in mime or mime == "application/vnd.ms-excel":
        try:
            import openpyxl, io  # type: ignore[import]
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in wb.worksheets:
                parts.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    row_str = "\t".join("" if v is None else str(v) for v in row)
                    if row_str.strip():
                        parts.append(row_str)
            wb.close()
            return "\n".join(parts)
        except ImportError:
            return "(XLSX extraction unavailable — install openpyxl)"
        except Exception as exc:
            return f"(XLSX extraction failed: {exc})"

    # Plain text / code
    if mime.startswith("text/") or ext in _TEXT_EXTENSIONS:
        return raw.decode("utf-8", errors="replace")

    # Unknown binary — try UTF-8, check printability
    try:
        text = raw.decode("utf-8", errors="strict")
        ratio = sum(1 for c in text if c.isprintable() or c in "\n\r\t") / max(len(text), 1)
        if ratio > 0.85:
            return text
    except UnicodeDecodeError:
        pass

    return "(binary file — no text content extractable)"


def _build_user_content(prompt: str, req: GenerateRequest, vision_supported: bool = True) -> str | list:
    """
    Build the user message content.
    - Images on vision models: sent as image_url blocks (OpenAI vision format).
    - Images on non-vision models: replaced with a text placeholder.
    - All other files (PDF, DOCX, XLSX, text, code, …): text extracted and injected as text blocks.
    """
    import base64

    attachments = req.inline_attachments if req.inline_attachments else []
    if not attachments:
        return prompt

    parts: list[dict] = []

    for att in attachments:
        mime = att.mime_type.lower()
        raw = base64.b64decode(att.data)

        if mime.startswith("image/"):
            if vision_supported:
                parts.append({
                    "type": "image_url",
                    "image_url": {"url": f"data:{att.mime_type};base64,{att.data}"},
                })
            else:
                parts.append({
                    "type": "text",
                    "text": f"[Image attached: {att.filename} — this model does not support image input]",
                })
        else:
            text = _extract_inline_text(raw, mime, att.filename)
            text = clamp_text_to_token_budget(
                text,
                INLINE_ATTACHMENT_TOKEN_BUDGET,
                marker="\n... [attached file content omitted from this request to fit the context budget]",
            )
            parts.append({"type": "text", "text": f"[File: {att.filename}]\n{text.strip()}"})

    parts.append({"type": "text", "text": prompt})

    # Collapse to plain string when no vision blocks present (maximises model compat)
    if not any(p.get("type") == "image_url" for p in parts):
        return "\n\n".join(
            p["text"] for p in parts if p.get("type") == "text" and p.get("text")
        )

    return parts


def build_messages(
    req: GenerateRequest,
    vision_supported: bool = True,
    *,
    model_context: int | None = None,
) -> list[dict]:
    """Build the full messages array including conversation history."""
    messages: list[dict] = []

    if req.system_message.strip():
        messages.append({"role": "system", "content": req.system_message.strip()})

    if req.continuation_text.strip():
        tail_messages = [
            {"role": "assistant", "content": req.continuation_text},
            {"role": "user", "content": CONTINUE_PROMPT},
        ]
    else:
        tail_messages = [
            {"role": "user", "content": _build_user_content(req.prompt, req, vision_supported)}
        ]

    reserved_tokens = req.max_tokens + sum(
        _content_token_estimate(message["content"]) + 4 for message in [*messages, *tail_messages]
    )
    history_budget = _resolve_history_budget(
        model_context=model_context,
        reserved_tokens=reserved_tokens,
    )
    history = select_recent_messages_by_token_budget(
        req.history,
        max_tokens=history_budget,
        max_messages=req.history_limit,
    )
    for h in history:
        messages.append({"role": h.role, "content": h.content})

    messages.extend(tail_messages)
    return messages


def _build_continuation_messages(base_messages: list[dict], partial_text: str) -> list[dict]:
    return [
        *base_messages,
        {"role": "assistant", "content": partial_text},
        {"role": "user", "content": CONTINUE_PROMPT},
    ]


def _merge_continuation_text(existing: str, addition: str) -> str:
    if not existing:
        return addition
    if not addition:
        return existing
    if addition in existing:
        return existing

    max_overlap = min(len(existing), len(addition), 200)
    for overlap in range(max_overlap, 0, -1):
        if existing[-overlap:] == addition[:overlap]:
            return existing + addition[overlap:]
    return existing + addition


def _looks_truncated(text: str) -> bool:
    stripped = text.rstrip()
    if not stripped:
        return False

    if stripped.count("```") % 2 == 1:
        return True

    if stripped.endswith(("...", "…", ":", ";", ",", "(", "[", "{", "-", "—", "–", "/", "\\")):
        return True

    if re.search(r"(^|\n)(?:[-*]|\d+\.)\s+[^\n]*$", stripped):
        return True

    return stripped[-1].isalnum()


def _should_auto_continue(
    *,
    finish_reason: str | None,
    text: str,
    max_tokens: int,
    completion_tokens: int | None = None,
) -> bool:
    if not text.strip():
        return False

    normalized_reason = (finish_reason or "").strip().lower()
    if normalized_reason in AUTO_CONTINUE_FINISH_REASONS:
        return True

    token_count = completion_tokens if completion_tokens is not None else estimate_tokens(text)
    near_limit = token_count >= max(32, int(max_tokens * 0.9))
    return near_limit and _looks_truncated(text)




async def real_generate(client: AsyncOpenAI, model_info: dict, req: GenerateRequest) -> GenerateResponse:
    """Call any OpenAI-compatible API with full conversation history."""
    base_messages = build_messages(
        req,
        vision_supported=model_info.get("vision", False),
        model_context=model_info.get("context"),
    )
    provider  = model_info["provider"]
    request_client = client.with_options(timeout=None)
    start = time.time()
    last_exc: Exception | None = None
    stream_started = False
    chosen_api_model: str | None = None
    full_text = ""
    total_prompt_tokens = 0
    total_completion_tokens = 0
    raw_passes: list[dict] = []
    last_finish_reason: str | None = None
    truncated = False
    provider_generation_id: str | None = None

    for api_model in _api_name_candidates(model_info):
        current_messages = base_messages
        full_text = ""
        total_prompt_tokens = 0
        total_completion_tokens = 0
        raw_passes = []
        last_finish_reason = None
        truncated = False
        provider_generation_id = None
        try:
            async with acquire_llm_capacity(api_model, provider):
                for pass_index in range(AUTO_CONTINUE_MAX_PASSES + 1):
                    effective_max_tokens = _effective_max_tokens(api_model, req.max_tokens)
                    create_kwargs: dict = {
                        "model":       api_model,
                        "messages":    current_messages,
                        "temperature": req.temperature,
                        "max_tokens":  effective_max_tokens,
                        "top_p":       req.top_p,
                    }

                    response = await request_client.chat.completions.create(**create_kwargs)
                    if provider_generation_id is None:
                        response_id = getattr(response, "id", None)
                        if isinstance(response_id, str) and response_id.strip():
                            provider_generation_id = response_id
                    text = response.choices[0].message.content or ""
                    full_text = _merge_continuation_text(full_text, text)
                    last_finish_reason = response.choices[0].finish_reason if response.choices else None
                    usage_data = response.usage

                    total_prompt_tokens += (
                        usage_data.prompt_tokens
                        if usage_data and usage_data.prompt_tokens is not None
                        else estimate_tokens(current_messages[-1]["content"] if current_messages else req.prompt)
                    )
                    total_completion_tokens += (
                        usage_data.completion_tokens
                        if usage_data and usage_data.completion_tokens is not None
                        else estimate_tokens(text)
                    )

                    try:
                        raw_passes.append(response.model_dump())
                    except Exception:
                        raw_passes.append({
                            "model": api_model,
                            "choices": [{
                                "message": {"content": text},
                                "finish_reason": last_finish_reason,
                            }],
                            "usage": {
                                "prompt_tokens": usage_data.prompt_tokens if usage_data else None,
                                "completion_tokens": usage_data.completion_tokens if usage_data else None,
                            },
                        })

                    completion_tokens = (
                        usage_data.completion_tokens
                        if usage_data and usage_data.completion_tokens is not None
                        else estimate_tokens(text)
                    )

                    should_continue = _should_auto_continue(
                        finish_reason=last_finish_reason,
                        text=text,
                        max_tokens=effective_max_tokens,
                        completion_tokens=completion_tokens,
                    )
                    if not should_continue:
                        break

                    if pass_index >= AUTO_CONTINUE_MAX_PASSES:
                        truncated = True
                        break

                    logger.info(
                        "[generate] auto_continue",
                        extra={
                            "provider": provider,
                            "model": api_model,
                            "pass_index": pass_index + 1,
                        },
                    )
                    current_messages = _build_continuation_messages(base_messages, full_text)
            chosen_api_model = api_model
            break
        except Exception as exc:
            last_exc = exc
            if _is_invalid_model_error(exc):
                logger.warning(
                    "[generate] invalid_model_alias_retry",
                    extra={"provider": provider, "model": api_model, "error": f"{type(exc).__name__}: {exc}"},
                )
                continue
            raise

    if chosen_api_model is None:
        if last_exc is not None:
            raise last_exc
        raise RuntimeError("no_valid_model_alias")

    latency = int((time.time() - start) * 1000)
    prompt_tok = total_prompt_tokens or estimate_tokens(req.prompt)
    completion_tok = total_completion_tokens or estimate_tokens(full_text)
    total_tokens = prompt_tok + completion_tok
    cost_usd = estimate_cost_usd(chosen_api_model, prompt_tok, completion_tok)

    usage = UsageStats(
        prompt_tokens=prompt_tok,
        completion_tokens=completion_tok,
        total_tokens=total_tokens,
        model=chosen_api_model,
        temperature=req.temperature,
        latency_ms=latency,
    )

    raw = {
        "model": chosen_api_model,
        "passes": raw_passes,
        "finish_reason": last_finish_reason,
        "continued_passes": max(0, len(raw_passes) - 1),
        "truncated": truncated,
        "provider_generation_id": provider_generation_id,
        "usage": {
            "prompt_tokens": usage.prompt_tokens,
            "completion_tokens": usage.completion_tokens,
            "total_tokens": usage.total_tokens,
            "cost_usd": cost_usd,
        },
    }

    return GenerateResponse(text=full_text, usage=usage, raw=raw, provider=provider)


async def real_generate_stream(client: AsyncOpenAI, model_info: dict, req: GenerateRequest):
    """Streaming generator - yields SSE-formatted chunks."""
    base_messages = build_messages(
        req,
        vision_supported=model_info.get("vision", False),
        model_context=model_info.get("context"),
    )
    request_client = client.with_options(timeout=None)
    api_model: str | None = None
    full_text = ""
    last_finish_reason: str | None = None
    truncated = False
    continued_passes = 0
    total_prompt_tokens = 0
    total_completion_tokens = 0
    total_tokens = 0
    provider_generation_id: str | None = None
    last_exc: Exception | None = None

    for candidate_api_model in _api_name_candidates(model_info):
        api_model = candidate_api_model
        full_text = ""
        last_finish_reason = None
        truncated = False
        continued_passes = 0
        total_prompt_tokens = 0
        total_completion_tokens = 0
        total_tokens = 0
        provider_generation_id = None
        try:
            async with acquire_llm_capacity(api_model, model_info["provider"]):
                current_messages = base_messages

                for pass_index in range(AUTO_CONTINUE_MAX_PASSES + 1):
                    effective_max_tokens = _effective_max_tokens(api_model, req.max_tokens)
                    create_kwargs: dict = {
                        "model": api_model,
                        "messages": current_messages,
                        "temperature": req.temperature,
                        "max_tokens": effective_max_tokens,
                        "top_p": req.top_p,
                        "stream": True,
                        "stream_options": {"include_usage": True},
                    }

                    pass_text = ""
                    last_finish_reason = None
                    pass_prompt_tokens: int | None = None
                    pass_completion_tokens: int | None = None
                    pass_total_tokens: int | None = None

                    async for chunk in await request_client.chat.completions.create(**create_kwargs):
                        if provider_generation_id is None:
                            chunk_id = getattr(chunk, "id", None)
                            if isinstance(chunk_id, str) and chunk_id.strip():
                                provider_generation_id = chunk_id
                        choice = chunk.choices[0] if chunk.choices else None
                        if choice and choice.finish_reason:
                            last_finish_reason = choice.finish_reason

                        usage_data = getattr(chunk, "usage", None)
                        if usage_data:
                            if usage_data.prompt_tokens is not None:
                                pass_prompt_tokens = usage_data.prompt_tokens
                            if usage_data.completion_tokens is not None:
                                pass_completion_tokens = usage_data.completion_tokens
                            if usage_data.total_tokens is not None:
                                pass_total_tokens = usage_data.total_tokens

                        delta = choice.delta.content if choice else None
                        if not delta:
                            continue

                        pass_text += delta
                        if pass_index == 0:
                            stream_started = True
                            full_text += delta
                            payload = json.dumps({"text": delta, "done": False}, ensure_ascii=False)
                            yield f"data: {payload}\n\n"

                    if pass_prompt_tokens is not None:
                        total_prompt_tokens += pass_prompt_tokens
                    if pass_completion_tokens is not None:
                        total_completion_tokens += pass_completion_tokens
                    if pass_total_tokens is not None:
                        total_tokens += pass_total_tokens

                    if pass_index > 0 and pass_text:
                        merged_text = _merge_continuation_text(full_text, pass_text)
                        addition = merged_text[len(full_text):]
                        full_text = merged_text

                        if addition:
                            for index in range(0, len(addition), 160):
                                payload = json.dumps(
                                    {"text": addition[index:index + 160], "done": False},
                                    ensure_ascii=False,
                                )
                                yield f"data: {payload}\n\n"

                    should_continue = _should_auto_continue(
                        finish_reason=last_finish_reason,
                        text=pass_text,
                        max_tokens=effective_max_tokens,
                        completion_tokens=pass_completion_tokens if pass_completion_tokens is not None else estimate_tokens(pass_text),
                    )
                    if not should_continue:
                        break

                    if pass_index >= AUTO_CONTINUE_MAX_PASSES:
                        truncated = True
                        break

                    logger.info(
                        "[stream] auto_continue",
                        extra={
                            "provider": model_info["provider"],
                            "model": api_model,
                            "pass_index": pass_index + 1,
                        },
                    )
                    continued_passes += 1
                    current_messages = _build_continuation_messages(base_messages, full_text)
            break
        except Exception as e:
            last_exc = e
            if _is_invalid_model_error(e):
                logger.warning(
                    "[stream] invalid_model_alias_retry",
                    extra={"provider": model_info["provider"], "model": api_model, "error": f"{type(e).__name__}: {e}"},
                )
                continue
            logger.error(f"[stream] Error: {e}")
            if isinstance(e, APIStatusError):
                raise RuntimeError(f"provider_http_{e.status_code}") from e
            raise

    if api_model is None or (last_exc is not None and not stream_started and not full_text):
        if last_exc is not None:
            logger.error(f"[stream] Error: {last_exc}")
            if isinstance(last_exc, APIStatusError):
                raise RuntimeError(f"provider_http_{last_exc.status_code}") from last_exc
            raise last_exc

    exact_total_tokens = total_tokens if total_tokens > 0 else None
    exact_prompt_tokens = total_prompt_tokens if total_prompt_tokens > 0 else None
    exact_completion_tokens = total_completion_tokens if total_completion_tokens > 0 else None
    if exact_completion_tokens is None and exact_total_tokens is not None and exact_prompt_tokens is not None:
        exact_completion_tokens = max(0, exact_total_tokens - exact_prompt_tokens)
    if exact_total_tokens is None and exact_prompt_tokens is not None and exact_completion_tokens is not None:
        exact_total_tokens = exact_prompt_tokens + exact_completion_tokens

    usage_payload = None
    if exact_total_tokens is not None:
        prompt_for_cost = exact_prompt_tokens or 0
        completion_for_cost = exact_completion_tokens
        if completion_for_cost is None:
            completion_for_cost = max(0, exact_total_tokens - prompt_for_cost)
        usage_payload = {
            "prompt_tokens": exact_prompt_tokens,
            "completion_tokens": exact_completion_tokens,
            "total_tokens": exact_total_tokens,
            "cost_usd": estimate_cost_usd(api_model, prompt_for_cost, completion_for_cost),
        }

    yield f"data: {json.dumps({'text': '', 'done': True, 'full_text': full_text, 'finish_reason': last_finish_reason, 'continued_passes': continued_passes, 'truncated': truncated, 'usage': usage_payload, 'provider_generation_id': provider_generation_id}, ensure_ascii=False)}\n\n"


def _has_real_api_keys() -> bool:
    return bool(OPENROUTER_API_KEY)


# Refine prompt with LLM
LEVEL_GUIDANCE = {
    1: """USER LEVEL: Beginner (L1)
- Use warm, plain language and avoid prompt-engineering jargon.
- Teach one useful habit at a time. Do not overwhelm the user.
- If the prompt is weak, fix the highest-impact missing piece first.
- Ask simple clarifying questions that a first-time user can answer quickly.
- The improved prompt should feel natural, clear, and not over-engineered.""",

    2: """USER LEVEL: Intermediate (L2)
- Be direct and practical.
- Focus on task specificity, context, output format, constraints, and success criteria.
- Show better structure without turning every prompt into a long template.
- Ask only the questions that remove real ambiguity.
- Explain the reasoning behind the edits in practical terms.""",

    3: """USER LEVEL: Advanced (L3)
- Be concise and peer-level.
- Look for hidden assumptions, missing constraints, vague evaluation criteria, edge cases, and underspecified output contracts.
- Ask only questions that materially improve precision or reduce failure risk.
- The improved prompt should be precise, minimal, and production-ready.
- Explain only the subtle, high-impact changes.""",
}

PROMPT_TUTOR_SYSTEM = """You are Nexa's Prompt Coach, an expert prompt-engineering assistant embedded in an AI chat product.
Your job is to help the user send a better prompt while teaching a small reusable principle.
You must understand the user's chat context, infer the task domain, and coach as if you are a competent professional in that domain.

{level_guidance}

## What to optimize
Evaluate the prompt using these criteria, but only mention issues that matter for this exact request:
1. Intent clarity — is the task/action unambiguous?
2. Context sufficiency — does the model have enough background, source material, constraints, and user goal?
3. Output contract — format, length, tone, audience, language, and structure.
4. Success criteria — what would make the answer useful or correct?
5. Constraints and risks — hard limits, exclusions, edge cases, safety, factual precision.
6. Domain fit — adapt to the topic without judging the user's domain expertise as their interface skill.
7. Prompt economy — avoid unnecessary roles, long templates, and generic filler.

## Scenario handling
- Weak or vague prompt: identify the main missing piece, ask 1-3 clarifying questions, and produce a stronger draft using only safe assumptions.
- Good prompt: do not invent major problems. Make a light edit for clarity, grammar, structure, or precision. gaps may be empty or contain one minor improvement.
- Excellent prompt: preserve most of it. Improve only wording, formatting, or small ambiguity. Say that it is already strong.
- Very long prompt: preserve the user's structure and intent. Do not compress away requirements. Improve organization, remove duplication only when obvious, and keep important details.
- Creative prompts: preserve voice and creative direction. Do not over-formalize.
- Coding/technical prompts: preserve exact identifiers, versions, constraints, errors, and code intent. Do not change the requested stack.
- If recent chat context makes the user's intent clear, use it. Do not ask a question that the chat history already answers.
- If the user supplied clarifying answers, integrate them into improved_prompt. In that second pass, opening_message should confirm the prompt is now sharper and should NOT repeat "what is missing" unless a critical unresolved ambiguity remains.

## Response rules
- Respond ONLY in {response_language}. Every field must be in this language.
- If {response_language} is Ukrainian, write the explanatory prose in natural Ukrainian. Do not write English coaching sentences such as "Your prompt", "No specification", "What is your..." or "Unclear...".
- English technical terms are allowed inside Ukrainian text when they are standard or user-provided terms: API, JSON, React, OAuth, REST, prompt, endpoint, debounce, stack traces, model names, code identifiers, product names, formulas, and quoted text.
- Output ONLY raw JSON. No markdown fences. No explanation outside JSON.
- opening_message: 1-2 short sentences. Say what you changed and why. Avoid generic praise.
- strengths: 0-2 specific items. Do not force strengths if they would be fake.
- gaps: 0-3 specific items. For second pass with clarifying answers, keep this empty unless something critical is still unresolved.
- clarifying_questions: 1-3 questions when answers would materially improve the prompt; 0 if the prompt is already clear enough. Each question must be answerable in one sentence. Do not ask broad or redundant questions.
- improved_prompt: a ready-to-send prompt. Preserve intent, facts, language, names, code, and constraints. Never make the prompt less specific.
- why_this_is_better: 1-3 concise items explaining concrete edits and their effect.

## Anti-patterns to avoid
- Do not suggest temperature, top-p, system prompts, variables, or any UI controls.
- Do not pad the improved_prompt with generic phrases like "As an expert..." unless role is genuinely needed.
- Do not repeat the gap diagnosis verbatim in why_this_is_better.
- Do not generate a worse prompt just to seem different.
- Do not make up facts that are not in the user prompt or chat history.
- Do not turn every prompt into a rigid checklist.
- Do not criticize the user. The tone should be calm, precise, and useful.

Output schema, strict, no extra keys:
{{
  "opening_message": "string",
  "strengths": ["string"],
  "gaps": ["string"],
  "clarifying_questions": [
    {{"id": "q1", "question": "string"}},
    {{"id": "q2", "question": "string"}},
    {{"id": "q3", "question": "string"}}
  ],
  "improved_prompt": "string",
  "why_this_is_better": ["string"]
}}"""


UKRAINIAN_MARKERS = (
    "\u0456", "\u0457", "\u0454", "\u0491",
    "\u0406", "\u0407", "\u0404", "\u0490",
)


def _detect_language(text: str) -> str:
    """Heuristic language detection for tutor responses."""
    if any(ch in text for ch in UKRAINIAN_MARKERS):
        return "uk"

    cyrillic_count = sum(1 for ch in text if "\u0400" <= ch <= "\u04FF")
    latin_count = sum(1 for ch in text if "a" <= ch.lower() <= "z")

    if cyrillic_count == 0:
        return "en"
    if latin_count == 0:
        return "uk"
    return "uk" if cyrillic_count >= max(1, int(latin_count * 0.25)) else "en"


def _is_ukrainian_user_facing_text(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return True
    if any(ch in stripped for ch in UKRAINIAN_MARKERS):
        return True

    cyrillic_count = sum(1 for ch in stripped if "\u0400" <= ch <= "\u04FF")
    latin_count = sum(1 for ch in stripped if "a" <= ch.lower() <= "z")
    if cyrillic_count == 0 and latin_count == 0:
        return True
    return cyrillic_count > 0 and cyrillic_count >= max(1, int(latin_count * 0.15))


def _response_matches_language(data: dict, lang: str) -> bool:
    if lang == "uk":
        fields: list[str] = [
            str(data.get("opening_message", "")),
            str(data.get("improved_prompt", "")),
            *[str(item) for item in data.get("strengths", [])],
            *[str(item) for item in data.get("gaps", [])],
            *[str(item) for item in data.get("why_this_is_better", [])],
            *[
                str(item.get("question", ""))
                for item in data.get("clarifying_questions", [])
                if isinstance(item, dict)
            ],
        ]
        return all(_is_ukrainian_user_facing_text(field) for field in fields)
    return True


def _language_label(lang: str) -> str:
    return "Ukrainian" if lang == "uk" else "English"


def _build_tutor_system_prompt(language: str, level: int) -> str:
    level_guidance = LEVEL_GUIDANCE.get(level, LEVEL_GUIDANCE[1])
    return PROMPT_TUTOR_SYSTEM.format(
        level_guidance=level_guidance,
        response_language=_language_label(language),
    )


def _build_tutor_user_message(
    prompt: str,
    clarification_answers: dict[str, str] | None = None,
    *,
    history: list[HistoryMessage] | None = None,
    history_limit: int = 30,
) -> str:
    parts: list[str] = []
    recent_history = select_recent_messages_by_token_budget(
        history or [],
        max_tokens=TUTOR_HISTORY_TOKEN_BUDGET,
        max_messages=history_limit,
    )
    if recent_history:
        history_text = "\n".join(f"{item.role}: {item.content}" for item in recent_history)
        parts.append(
            "Recent chat context. Use it to infer the task and avoid redundant questions. "
            "Newest relevant messages are retained; oldest messages may be omitted:\n"
            f"{history_text}"
        )

    prompt_language = _language_label(_detect_language(prompt))
    parts.append(f'User prompt: "{prompt}"')
    parts.append(
        f"Detected user prompt language: {prompt_language}. "
        f"All user-facing JSON values must be written in {prompt_language}, regardless of the app UI language."
    )
    if clarification_answers:
        answers_text = "\n".join(
            f"- {qid}: {answer}" for qid, answer in clarification_answers.items() if answer.strip()
        )
        if answers_text:
            parts.append(f"\nUser answered clarifying questions:\n{answers_text}")
            parts.append(
                "\nThis is the second pass. Integrate the answers into improved_prompt. "
                "Do not repeat missing-information notes unless a critical ambiguity remains."
            )
    return "\n".join(parts)


def _normalize_tutor_response(parsed: dict) -> dict:
    """Validate and normalize the parsed JSON into the expected response shape."""
    opening = str(parsed.get("opening_message", "")).strip()
    strengths = [str(s).strip() for s in parsed.get("strengths", []) if str(s).strip()][:2]
    gaps = [str(g).strip() for g in parsed.get("gaps", []) if str(g).strip()][:3]

    raw_questions = parsed.get("clarifying_questions", [])
    questions: list[dict] = []
    for i, q in enumerate(raw_questions):
        if isinstance(q, dict) and q.get("question"):
            questions.append({
                "id": str(q.get("id", f"q{i+1}")),
                "question": str(q["question"]).strip(),
            })
        elif isinstance(q, str) and q.strip():
            questions.append({"id": f"q{i+1}", "question": q.strip()})
    questions = questions[:3]  # keep the UI focused and cognitively light

    improved = str(parsed.get("improved_prompt", "")).strip()
    why_better = [str(w).strip() for w in parsed.get("why_this_is_better", []) if str(w).strip()][:3]

    if not improved or len(improved) < 10:
        raise ValueError("low_quality_response: improved_prompt too short")

    return {
        "opening_message": opening,
        "strengths": strengths,
        "gaps": gaps,
        "clarifying_questions": questions,
        "improved_prompt": improved,
        "why_this_is_better": why_better,
        "next_step": "",
    }


def _iter_tutor_candidates() -> list[tuple[AsyncOpenAI, dict]]:
    """
    Ordered list of (client, model_info) for tutor fallback chain.
    Strongest reliable reasoning models first, then cheaper fast models.
    """
    preferred_models = [
        "claude-sonnet-4-5",
        "gemini-2.5-pro",
        "gpt-4o",
        "o4-mini",
        "claude-haiku-4-5",
        "gemini-2.5-flash",
        "gemini-2.0-flash",
        "gpt-4o-mini",
    ]
    seen: set[str] = set()
    candidates: list[tuple[AsyncOpenAI, dict]] = []

    for model_id in preferred_models:
        info = AVAILABLE_MODELS.get(model_id)
        if not info:
            continue
        client = clients.get(info["provider"])
        if not client:
            continue
        candidates.append((client, info))
        seen.add(model_id)

    # Append any remaining available models not already in the list
    for model_id, info in AVAILABLE_MODELS.items():
        if model_id in seen:
            continue
        client = clients.get(info["provider"])
        if not client:
            continue
        candidates.append((client, info))

    return candidates


async def refine_prompt_with_llm(
    prompt: str,
    *,
    language: str | None = None,
    level: int | None = None,
    clarification_answers: dict[str, str] | None = None,
    history: list[HistoryMessage] | None = None,
    history_limit: int = 30,
) -> dict:
    lang = language if language in ("en", "uk") else _detect_language(prompt)
    lvl = level if level in (1, 2, 3) else 1

    candidates = _iter_tutor_candidates()
    if not candidates:
        raise RuntimeError("no_client")

    system_prompt = _build_tutor_system_prompt(lang, lvl)
    user_message = _build_tutor_user_message(
        prompt,
        clarification_answers,
        history=history,
        history_limit=history_limit,
    )
    last_exc: Exception | None = None

    for client, model_info in candidates:
        provider = model_info["provider"]
        request_client = client.with_options(timeout=None)
        for api_model in _api_name_candidates(model_info):
            create_kwargs: dict = {
                "model": api_model,
                "messages": [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_message},
                ],
                "temperature": 0.35,
                "max_tokens": 1800,
            }

            # OpenRouter supports json_object for Claude and GPT models
            if "anthropic/" in api_model or "openai/" in api_model:
                create_kwargs["response_format"] = {"type": "json_object"}

            try:
                async with acquire_llm_capacity(api_model, provider):
                    response = await request_client.chat.completions.create(**create_kwargs)

                content = (response.choices[0].message.content or "").strip()
                content = re.sub(r"^```(?:json)?\s*", "", content)
                content = re.sub(r"\s*```$", "", content)
                json_match = re.search(r'\{[\s\S]*\}', content)
                if json_match:
                    content = json_match.group(0)

                parsed = json.loads(content.strip())
                normalized = _normalize_tutor_response(parsed)
                if not _response_matches_language(normalized, lang):
                    raise ValueError(f"wrong_tutor_language:{lang}")
                has_clarification_answers = bool(
                    clarification_answers
                    and any(value.strip() for value in clarification_answers.values())
                )
                if has_clarification_answers:
                    normalized["gaps"] = []
                    normalized["clarifying_questions"] = []
                return normalized
            except Exception as exc:
                last_exc = exc
                if _is_invalid_model_error(exc):
                    logger.warning(
                        "[refine] tutor_invalid_model_alias_retry",
                        extra={"provider": provider, "model": api_model, "error": f"{type(exc).__name__}: {exc}"},
                    )
                    continue
                logger.warning(
                    "[refine] tutor_provider_failed",
                    extra={
                        "provider": provider,
                        "model": api_model,
                        "error": f"{type(exc).__name__}: {exc}",
                    },
                )
                break

    if last_exc is not None:
        raise last_exc
    raise RuntimeError("no_client")


PROMPT_SUGGESTIONS_SYSTEM = """You generate short, high-quality prompt suggestions for an AI chat app.

Rules:
- Return ONLY raw JSON with key "suggestions".
- Generate 4 suggestions.
- Each suggestion must be one sentence, 6-16 words, ready to click and send.
- When recent context exists, infer the user's recurring topics and generate suggestions connected to those topics.
- Do not simply repeat old messages. Propose useful next prompts the user may want now.
- If context is weak or this is a new account, create diverse cold-start prompts for different user types: learning, coding, writing, planning, research, and work.
- Do not mention UI levels, settings, or the app.
- Use the requested interface language for all user-facing text.
- English technical terms are allowed inside Ukrainian suggestions when they are natural terms, but the sentence must remain Ukrainian.
- Keep suggestions specific, practical, and safe. Avoid generic prompts like "Tell me something interesting".

Schema:
{"suggestions":["string","string","string","string"]}"""


PROMPT_SUGGESTION_FALLBACKS: dict[str, list[str]] = {
    "en": [
        "Explain a difficult concept with a simple example",
        "Create a study plan for learning a new skill",
        "Review this idea and suggest stronger alternatives",
        "Write a clear email for a professional situation",
        "Compare two approaches and recommend the better one",
        "Turn messy notes into a structured action plan",
        "Help me debug a problem step by step",
        "Summarize a complex topic for a beginner",
    ],
    "uk": [
        "Поясни складну тему на простому прикладі",
        "Склади план навчання для нової навички",
        "Оціни ідею та запропонуй сильніші варіанти",
        "Напиши зрозумілий лист для робочої ситуації",
        "Порівняй два підходи й порадь кращий",
        "Перетвори нотатки на структурований план дій",
        "Допоможи покроково розібрати технічну проблему",
        "Поясни складну тему для початківця",
    ],
}


def fallback_prompt_suggestions(language: str | None = None) -> list[str]:
    lang = language if language in ("en", "uk") else "en"
    suggestions = PROMPT_SUGGESTION_FALLBACKS[lang][:]
    random.shuffle(suggestions)
    return suggestions[:4]


def _build_prompt_suggestions_user_message(
    *,
    language: str,
    level: int,
    history: list[HistoryMessage] | None = None,
) -> str:
    recent_history = select_recent_messages_by_token_budget(
        history or [],
        max_tokens=1600,
        max_messages=12,
    )
    history_text = "\n".join(f"{item.role}: {item.content}" for item in recent_history)
    if not history_text.strip():
        history_text = "No useful previous context yet. Use cold-start diversity."

    return (
        f"Language: {_language_label(language)}\n"
        f"User interface level: {level}\n"
        "Recent conversation context:\n"
        f"{history_text}\n\n"
        "Generate prompt suggestions now. Return only the JSON object."
    )


def _normalize_prompt_suggestions(parsed: dict, lang: str) -> list[str]:
    raw = parsed.get("suggestions", [])
    suggestions = [str(item).strip() for item in raw if str(item).strip()]
    suggestions = [item for item in suggestions if 6 <= len(item) <= 140]
    if len(suggestions) < 3:
        raise ValueError("low_quality_prompt_suggestions")
    if lang == "uk" and not all(_is_ukrainian_user_facing_text(item) for item in suggestions):
        raise ValueError("wrong_suggestion_language")
    return suggestions[:4]


def _iter_prompt_suggestion_candidates() -> list[tuple[AsyncOpenAI, dict]]:
    preferred_models = [
        "claude-haiku-4-5",
        "gemini-2.5-flash",
        "gemini-2.0-flash",
        "gpt-4o-mini",
        "claude-sonnet-4-5",
        "gemini-2.5-pro",
        "gpt-4o",
    ]
    seen: set[str] = set()
    candidates: list[tuple[AsyncOpenAI, dict]] = []
    for model_id in preferred_models:
        info = AVAILABLE_MODELS.get(model_id)
        if not info:
            continue
        client = clients.get(info["provider"])
        if not client:
            continue
        candidates.append((client, info))
        seen.add(model_id)
    for model_id, info in AVAILABLE_MODELS.items():
        if model_id in seen:
            continue
        client = clients.get(info["provider"])
        if client:
            candidates.append((client, info))
    return candidates


async def generate_prompt_suggestions_with_llm(
    *,
    language: str | None = None,
    level: int | None = None,
    history: list[HistoryMessage] | None = None,
) -> list[str]:
    history_text = "\n".join(item.content for item in (history or []) if item.content.strip())
    lang = language if language in ("en", "uk") else _detect_language(history_text)
    lvl = level if level in (1, 2, 3) else 1
    candidates = _iter_prompt_suggestion_candidates()
    if not candidates:
        raise RuntimeError("no_client")

    user_message = _build_prompt_suggestions_user_message(language=lang, level=lvl, history=history)
    last_exc: Exception | None = None
    for client, model_info in candidates:
        provider = model_info["provider"]
        request_client = client.with_options(timeout=20)
        for api_model in _api_name_candidates(model_info):
            create_kwargs: dict = {
                "model": api_model,
                "messages": [
                    {"role": "system", "content": PROMPT_SUGGESTIONS_SYSTEM},
                    {"role": "user", "content": user_message},
                ],
                "temperature": 0.7,
                "max_tokens": 420,
            }
            if "anthropic/" in api_model or "openai/" in api_model:
                create_kwargs["response_format"] = {"type": "json_object"}
            try:
                async with acquire_llm_capacity(api_model, provider):
                    response = await request_client.chat.completions.create(**create_kwargs)
                content = (response.choices[0].message.content or "").strip()
                content = re.sub(r"^```(?:json)?\s*", "", content)
                content = re.sub(r"\s*```$", "", content)
                json_match = re.search(r'\{[\s\S]*\}', content)
                if json_match:
                    content = json_match.group(0)
                return _normalize_prompt_suggestions(json.loads(content), lang)
            except Exception as exc:
                last_exc = exc
                if _is_invalid_model_error(exc):
                    continue
                logger.warning(
                    "[suggestions] provider_failed",
                    extra={"provider": provider, "model": api_model, "error": f"{type(exc).__name__}: {exc}"},
                )
                break
    if last_exc is not None:
        raise last_exc
    raise RuntimeError("no_client")
